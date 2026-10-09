"""
PPTX Parser for extracting slide text and speaker notes from presentations.

Uses python-pptx. Text and notes only — images, charts, SmartArt, animations,
and non-trivial tables are out of scope (Stage 1, Milestone 8).
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

from pptx import Presentation

# Cap large decks; even 100+ slide decks are usually small in text, but a cap
# is prudent (project-owner decision, spec §7 Q3).
DEFAULT_MAX_SLIDES = 200

# Safety net so a text-heavy deck cannot blow the agent's context window.
MAX_CHARS = 200_000

_TRUNCATION_MARKER = "\n\n[... truncated]"


@dataclass
class PptxSlide:
    """Represents extracted content from a single slide."""
    slide_number: int  # 1-indexed
    title: Optional[str]
    body_text: List[str] = field(default_factory=list)
    notes: str = ""

    def to_markdown(self) -> str:
        """Convert slide content to markdown format."""
        heading = f"## Slide {self.slide_number}"
        if self.title:
            heading += f": {self.title}"
        parts = [heading]

        for block in self.body_text:
            if block.strip():
                parts.append(block.strip())

        if self.notes.strip():
            parts.append(f"_Notes:_ {self.notes.strip()}")

        return "\n\n".join(parts)


@dataclass
class PptxParseResult:
    """Result of parsing a PPTX presentation."""
    filename: str
    slides: List[PptxSlide] = field(default_factory=list)
    total_slides: int = 0
    error: Optional[str] = None

    @property
    def slide_count(self) -> int:
        """Number of slides actually extracted."""
        return len(self.slides)

    def to_markdown(self) -> str:
        """Convert all extracted slides to markdown."""
        if self.error:
            return f"[Error parsing PPTX: {self.error}]"

        parts = [
            f"=== PPTX: {self.filename} "
            f"({self.total_slides} slide{'s' if self.total_slides != 1 else ''}) ==="
        ]
        for slide in self.slides:
            parts.append(slide.to_markdown())

        if self.slide_count < self.total_slides:
            parts.append(
                f"[... truncated, showing first {self.slide_count} "
                f"of {self.total_slides} slides]"
            )

        markdown = "\n\n".join(parts)
        if len(markdown) > MAX_CHARS:
            markdown = markdown[:MAX_CHARS] + _TRUNCATION_MARKER
        return markdown

    def get_token_estimate(self, chars_per_token: float = 4.0) -> int:
        """Estimate token count based on character count."""
        return int(len(self.to_markdown()) / chars_per_token)


class PptxParser:
    """Parser for extracting slide text and notes from PPTX presentations."""

    DEFAULT_MAX_SLIDES = DEFAULT_MAX_SLIDES

    def parse(self, file_path: str, max_slides: Optional[int] = None) -> PptxParseResult:
        """Parse a PPTX file and extract slide text + speaker notes.

        Never raises for malformed input — returns a result with ``error`` set.
        """
        path = Path(file_path)
        filename = path.name

        if max_slides is None:
            max_slides = self.DEFAULT_MAX_SLIDES

        try:
            presentation = Presentation(str(path))
        except Exception as e:  # noqa: BLE001 - surface parse failures as error
            return PptxParseResult(filename=filename, slides=[], error=str(e))

        all_slides = list(presentation.slides)
        total_slides = len(all_slides)

        slides: List[PptxSlide] = []
        for idx, slide in enumerate(all_slides[:max_slides], start=1):
            slides.append(self._extract_slide(slide, idx))

        return PptxParseResult(
            filename=filename,
            slides=slides,
            total_slides=total_slides,
        )

    def _extract_slide(self, slide, slide_number: int) -> PptxSlide:
        """Extract title, body text, and notes from a single slide."""
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:  # noqa: BLE001 - some layouts have no title placeholder
            title_shape = None

        title = None
        title_id = None
        if title_shape is not None and title_shape.has_text_frame:
            title = title_shape.text.strip() or None
            # python-pptx returns a fresh wrapper per access, so identity (``is``)
            # comparison fails; match the title by its stable shape id instead.
            title_id = title_shape.shape_id

        body_text: List[str] = []
        for shape in slide.shapes:
            if title_id is not None and shape.shape_id == title_id:
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text
            if text and text.strip():
                body_text.append(text)

        notes = ""
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None and notes_frame.text:
                notes = notes_frame.text

        return PptxSlide(
            slide_number=slide_number,
            title=title,
            body_text=body_text,
            notes=notes,
        )
